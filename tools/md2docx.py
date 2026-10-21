"""Markdown → Word(.docx) 转换器
用法：python tools/md2docx.py 输入.md [输出.docx]

⚠️ 2026-10-20 用户定：「发出的产物最好不要是 .md，要么直接转为 word 文档，
   除非实在不行再用 md」⇒ 交付前遇到 .md 就先过这一道 ✓

支持：标题 1-3 级 / 有序无序列表 / 粗体 / 斜体 / 行内代码 / 代码块 /
     表格 / 引用 / 分隔线 / 普通段落。中文正文用微软雅黑。
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CN_FONT = '微软雅黑'
MONO = 'Consolas'


def _set_font(run, name=CN_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def _inline(par, text):
    """把 **粗体** / *斜体* / `代码` 拆成 run"""
    for seg in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**') and len(seg) > 4:
            _set_font(par.add_run(seg[2:-2]), bold=True)
        elif seg.startswith('*') and seg.endswith('*') and len(seg) > 2:
            _set_font(par.add_run(seg[1:-1]), italic=True)
        elif seg.startswith('`') and seg.endswith('`') and len(seg) > 2:
            r = par.add_run(seg[1:-1])
            _set_font(r, name=MONO, size=10)
            r.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
        else:
            _set_font(par.add_run(seg))


def _is_table_row(line):
    return line.strip().startswith('|') and line.strip().endswith('|')


def _split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def convert(src, dst=None):
    src = os.path.abspath(src)
    if not os.path.exists(src):
        raise FileNotFoundError(src)
    dst = dst or os.path.splitext(src)[0] + '.docx'
    lines = open(src, encoding='utf-8').read().splitlines()

    doc = Document()
    # 正文默认字体 / 字号
    st = doc.styles['Normal']
    st.font.name = CN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()

        # 空行
        if not s:
            i += 1
            continue

        # 代码块
        if s.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            for cl in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(cl if cl else ' ')
                _set_font(r, name=MONO, size=9.5)
            continue

        # 表格
        if _is_table_row(s):
            rows = []
            while i < n and _is_table_row(lines[i]):
                rows.append(_split_row(lines[i]))
                i += 1
            # 去掉 |---|---| 那条分隔
            rows = [r for r in rows if not all(re.fullmatch(r':?-{2,}:?', c or '-') for c in r)]
            if rows:
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=cols)
                t.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        cell = t.cell(ri, ci)
                        cell.text = ''
                        par = cell.paragraphs[0]
                        _inline(par, row[ci] if ci < len(row) else '')
                        if ri == 0:
                            for r0 in par.runs:
                                r0.bold = True
            continue

        # 分隔线
        if re.fullmatch(r'(-{3,}|\*{3,}|_{3,})', s):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pr = p._p.get_or_add_pPr()
            bd = pr.makeelement(qn('w:pBdr'), {})
            bottom = bd.makeelement(qn('w:bottom'), {})
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '4')
            bottom.set(qn('w:color'), 'CCCCCC')
            bd.append(bottom)
            pr.append(bd)
            i += 1
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            lv = min(3, len(m.group(1)))
            p = doc.add_heading('', level=lv)
            _inline(p, m.group(2))
            for r in p.runs:
                r.font.name = CN_FONT
                r._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
                r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            i += 1
            continue

        # 引用
        if s.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _inline(p, s.lstrip('>').strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # 列表（有序 / 无序，支持缩进的二级）
        m = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', raw)
        if m:
            indent = len(m.group(1).expandtabs(2))
            ordered = not m.group(2) in '-*+'
            style = 'List Number' if ordered else 'List Bullet'
            if indent >= 2:
                style = 'List Number 2' if ordered else 'List Bullet 2'
            p = doc.add_paragraph(style=style)
            _inline(p, m.group(3))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _inline(p, s)
        i += 1

    doc.save(dst)
    return dst


def main():
    if len(sys.argv) < 2:
        print('用法: python tools/md2docx.py 输入.md [输出.docx]')
        return 2
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else None
    out = convert(src, dst)
    print('OK ' + out + ' ' + str(os.path.getsize(out) // 1024) + 'KB')
    return 0


if __name__ == '__main__':
    sys.exit(main())
